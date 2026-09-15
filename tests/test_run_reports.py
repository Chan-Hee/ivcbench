"""Raw-run reports must stay separate from the curated scientific submission."""

from pathlib import Path
import ast

import numpy as np
import pandas as pd
from docx import Document

from ivcbench.clusters.c5 import cross_celltype_loct
from ivcbench.report.docx_export import build_docx
from ivcbench.report.draft import run_report, write_draft
from ivcbench.report.methods import cluster_methods, framework_methods, write_methods


def result_rows() -> pd.DataFrame:
    return pd.DataFrame(
        [
            dict(
                split="C5_cross_celltype_loct",
                baseline="example",
                action="run_headline",
                ran=True,
                pearson_delta=0.123456,
                e_distance=0.5,
                aucell_program_corr=np.nan,
                n_train=40,
                n_test=10,
                n_test_strata=2,
            ),
            dict(
                split="C5_cross_celltype_loct",
                baseline="failed-example",
                action="failed",
                ran=False,
                pearson_delta=np.nan,
            ),
        ]
    )


def test_report_uses_only_completed_rows_and_explicit_missing_scores():
    report = run_report("C5", result_rows(), "example source")
    assert "1 of 2" in report and "0.123" in report
    assert "| NA |" in report and "failed-example" not in report
    assert "not the curated submission panel" in report
    assert "Runtime applicability flags do not certify" in report
    for stale_claim in ("OnePager", "Figure 7", "leak-proof", "principled value"):
        assert stale_claim not in report
    figure_source = (
        Path(__file__).resolve().parents[1] / "src/ivcbench/report/figures.py"
    )
    figure_text = "\n".join(
        node.value
        for node in ast.walk(ast.parse(figure_source.read_text()))
        if isinstance(node, ast.Constant) and isinstance(node.value, str)
    )
    for stale_claim in (
        "flat — no Tanimoto difficulty effect",
        "conditioned models stay",
        "simple-only floor",
        "true leave-one-gene-out",
        "OnePager",
    ):
        assert stale_claim not in figure_text


def test_empty_and_synthetic_reports_have_no_invented_results():
    report = run_report("C3", pd.DataFrame(), "synthetic-test")
    assert "Synthetic fixture" in report and "0 of 0" in report
    assert "No completed result rows" in report


def test_methods_report_recorded_splits_not_template_datasets():
    spec = cross_celltype_loct("NK")
    rows = result_rows().assign(split=spec.name)
    report = cluster_methods("C5", rows, [spec], {"data_source": "synthetic-test"})
    assert "n_train = 40" in report and "n_test = 10" in report
    assert "Synthetic fixture" in report and spec.name in report
    shared = framework_methods({"seeds": [4], "packages": {"numpy": "test"}})
    assert "Seeds requested: [4]" in shared and "numpy test" in shared
    assert "membership checks only" in shared
    assert "Undefined program correlations remain missing" in shared
    assert "~70" not in report and "GSE279945" not in report


def test_report_writers_and_docx_keep_compatibility_paths(tmp_path: Path):
    draft = write_draft("C5", result_rows(), "example", tmp_path / "draft_C5.md")
    cluster, shared = write_methods(
        "C5", result_rows(), [], {}, tmp_path / "notes.md", tmp_path / "shared.md"
    )
    assert cluster.is_file() and shared.is_file()
    output = build_docx("C5", draft.read_text(), None, tmp_path / "draft_C5.docx")
    doc = Document(output)
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 2
    assert doc.tables[0].cell(1, 5).text == "NA"


def test_methods_can_describe_a_failed_run_without_sizes():
    frame = pd.DataFrame([dict(split="example", ran=False)])
    report = cluster_methods("C5", frame, [cross_celltype_loct("NK")], {})
    assert "no completed size record" in report


def test_docx_renders_code_without_backticks(tmp_path: Path):
    path = build_docx(
        "C1", "Use `results_raw.csv` for raw rows.", None, tmp_path / "code.docx"
    )
    paragraph = Document(path).paragraphs[0]
    assert paragraph.text == "Use results_raw.csv for raw rows."
    code = next(run for run in paragraph.runs if run.text == "results_raw.csv")
    assert code.font.name == "Courier New"


def test_docx_keeps_escaped_pipe_in_a_single_table_cell(tmp_path: Path):
    rows = result_rows()
    rows.loc[0, "split"] = "donor=a|lineage=B"
    report = run_report("C1", rows, "example")
    path = build_docx("C1", report, None, tmp_path / "table.docx")
    table = Document(path).tables[0]
    assert len(table.columns) == 6
    assert table.cell(1, 0).text == "donor=a|lineage=B"
